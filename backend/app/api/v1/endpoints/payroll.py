"""HR, payroll & leave endpoints. Admin manages everything; teachers get a self-service
"/payroll/my/..." surface (own leaves, balances and approved/paid payslips)."""

import io
import uuid
from datetime import date

from fastapi import APIRouter, Depends, Query, Response, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.core.dependencies import require_role
from app.core.payroll_pdf import SHEET_COLUMNS, period_label, render_payroll_sheet, render_payslip, sheet_rows
from app.db.session import get_db
from app.models.payroll import LeaveRequestStatus, Payslip, PayrollRun, SalaryAdvanceStatus
from app.models.user import RoleEnum, User
from app.repositories.tenant_repo import TenantRepository
from app.schemas.payroll import (
    DepartmentCreate,
    DepartmentOut,
    DepartmentUpdate,
    DesignationCreate,
    DesignationOut,
    DesignationUpdate,
    EmployeeOut,
    EmployeeProfileOut,
    EmployeeProfileUpsert,
    EmployeeType,
    LeaveBalanceOut,
    LeaveDecision,
    LeaveRequestCreate,
    LeaveRequestOut,
    LeaveTypeCreate,
    LeaveTypeOut,
    LeaveTypeUpdate,
    MyLeaveRequestCreate,
    PayrollRunCreate,
    PayrollRunDetail,
    PayrollRunOut,
    PayslipAdjustmentCreate,
    PayslipOut,
    SalaryAdvanceCreate,
    SalaryAdvanceOut,
    SalaryStructureCreate,
    SalaryStructureOut,
)
from app.services.payroll_hr_service import PayrollHrService, employee_key_of
from app.services.payroll_leave_service import PayrollLeaveService
from app.services.payroll_run_service import PayrollRunService

router = APIRouter(prefix="/payroll", tags=["payroll"])

admin_only = require_role(RoleEnum.ADMIN)
teacher_only = require_role(RoleEnum.TEACHER)


def _tenant_name(db: Session, tenant_id: uuid.UUID) -> str:
    tenant = TenantRepository(db).get_by_id(tenant_id)
    return tenant.name if tenant else "School"


def _pdf(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter([content]), media_type="application/pdf",
        headers={"Content-Disposition": f"inline; filename={filename}"},
    )


def _payslip_pdf(db: Session, tenant_id: uuid.UUID, payslip: Payslip, run: PayrollRun) -> StreamingResponse:
    service = PayrollRunService(db)
    data = service.payslip_out(payslip, run).model_dump(mode="json")
    employee_type, employee_id = employee_key_of(payslip)
    hr = PayrollHrService(db)
    profile = hr.profiles.get_for_employee(tenant_id, *((employee_id, None) if employee_type == "teacher" else (None, employee_id)))
    if profile is not None:
        data["bank_name"] = profile.bank_name
        data["bank_account_no"] = profile.bank_account_no
        if profile.department_id is not None:
            dept = hr.departments.get_by_id(tenant_id, profile.department_id)
            data["department"] = dept.name if dept else None
    pdf = render_payslip(tenant_name=_tenant_name(db, tenant_id), payslip=data)
    return _pdf(pdf, f"payslip-{run.period_year}-{run.period_month:02d}.pdf")


# --- employees & HR profiles ---


@router.get("/employees", response_model=list[EmployeeOut])
def list_employees(
    status_filter: str | None = Query(default="active", alias="status", pattern="^(active|inactive|all)$"),
    q: str | None = Query(default=None),
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
) -> list[EmployeeOut]:
    return PayrollHrService(db).employees_view(current_user.tenant_id, status_filter, q)


@router.get("/employees/{employee_type}/{employee_id}/profile", response_model=EmployeeProfileOut | None)
def get_employee_profile(
    employee_type: EmployeeType,
    employee_id: uuid.UUID,
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollHrService(db).get_profile(current_user.tenant_id, employee_type, employee_id)


@router.put("/employees/{employee_type}/{employee_id}/profile", response_model=EmployeeProfileOut)
def upsert_employee_profile(
    employee_type: EmployeeType,
    employee_id: uuid.UUID,
    payload: EmployeeProfileUpsert,
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollHrService(db).upsert_profile(current_user.tenant_id, employee_type, employee_id, payload)


# --- departments & designations ---


@router.get("/departments", response_model=list[DepartmentOut])
def list_departments(current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollHrService(db).list_departments(current_user.tenant_id)


@router.post("/departments", response_model=DepartmentOut, status_code=status.HTTP_201_CREATED)
def create_department(payload: DepartmentCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollHrService(db).create_department(current_user.tenant_id, payload)


@router.patch("/departments/{department_id}", response_model=DepartmentOut)
def update_department(
    department_id: uuid.UUID, payload: DepartmentUpdate,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    return PayrollHrService(db).update_department(current_user.tenant_id, department_id, payload)


@router.delete("/departments/{department_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_department(department_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    PayrollHrService(db).delete_department(current_user.tenant_id, department_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/designations", response_model=list[DesignationOut])
def list_designations(current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollHrService(db).list_designations(current_user.tenant_id)


@router.post("/designations", response_model=DesignationOut, status_code=status.HTTP_201_CREATED)
def create_designation(payload: DesignationCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollHrService(db).create_designation(current_user.tenant_id, payload)


@router.patch("/designations/{designation_id}", response_model=DesignationOut)
def update_designation(
    designation_id: uuid.UUID, payload: DesignationUpdate,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    return PayrollHrService(db).update_designation(current_user.tenant_id, designation_id, payload)


@router.delete("/designations/{designation_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_designation(designation_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    PayrollHrService(db).delete_designation(current_user.tenant_id, designation_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# --- salary structures ---


@router.get("/salary-structures", response_model=list[SalaryStructureOut])
def list_salary_structures(
    employee_type: EmployeeType | None = Query(default=None),
    employee_id: uuid.UUID | None = Query(default=None),
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollHrService(db).list_structures(current_user.tenant_id, employee_type, employee_id)


@router.post("/salary-structures", response_model=SalaryStructureOut, status_code=status.HTTP_201_CREATED)
def create_salary_structure(
    payload: SalaryStructureCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)
):
    return PayrollHrService(db).create_structure(current_user.tenant_id, payload)


@router.delete("/salary-structures/{structure_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_salary_structure(structure_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    PayrollHrService(db).delete_structure(current_user.tenant_id, structure_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# --- leave types ---


@router.get("/leave-types", response_model=list[LeaveTypeOut])
def list_leave_types(
    active_only: bool = Query(default=False),
    current_user: User = Depends(require_role(RoleEnum.ADMIN, RoleEnum.TEACHER)),
    db: Session = Depends(get_db),
):
    only_active = active_only or current_user.role == RoleEnum.TEACHER
    return PayrollLeaveService(db).list_types(current_user.tenant_id, only_active)


@router.post("/leave-types", response_model=LeaveTypeOut, status_code=status.HTTP_201_CREATED)
def create_leave_type(payload: LeaveTypeCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollLeaveService(db).create_type(current_user.tenant_id, payload)


@router.patch("/leave-types/{leave_type_id}", response_model=LeaveTypeOut)
def update_leave_type(
    leave_type_id: uuid.UUID, payload: LeaveTypeUpdate,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    return PayrollLeaveService(db).update_type(current_user.tenant_id, leave_type_id, payload)


# --- leave requests (admin) ---


@router.get("/leaves", response_model=list[LeaveRequestOut])
def list_leave_requests(
    status_filter: LeaveRequestStatus | None = Query(default=None, alias="status"),
    employee_type: EmployeeType | None = Query(default=None),
    employee_id: uuid.UUID | None = Query(default=None),
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollLeaveService(db).list_requests(current_user.tenant_id, status_filter, employee_type, employee_id)


@router.post("/leaves", response_model=LeaveRequestOut, status_code=status.HTTP_201_CREATED)
def create_leave_request(payload: LeaveRequestCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollLeaveService(db)
    return service.to_out(service.apply_for_employee(current_user.tenant_id, payload, current_user.id))


@router.post("/leaves/{request_id}/approve", response_model=LeaveRequestOut)
def approve_leave_request(
    request_id: uuid.UUID, payload: LeaveDecision | None = None,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    service = PayrollLeaveService(db)
    return service.to_out(service.approve(current_user.tenant_id, request_id, current_user.id, payload.note if payload else None))


@router.post("/leaves/{request_id}/reject", response_model=LeaveRequestOut)
def reject_leave_request(
    request_id: uuid.UUID, payload: LeaveDecision | None = None,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    service = PayrollLeaveService(db)
    return service.to_out(service.reject(current_user.tenant_id, request_id, current_user.id, payload.note if payload else None))


@router.get("/leave-balances", response_model=list[LeaveBalanceOut])
def get_leave_balances(
    employee_type: EmployeeType,
    employee_id: uuid.UUID,
    year: int | None = Query(default=None, ge=2000, le=2100),
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollLeaveService(db).balances_for(current_user.tenant_id, employee_type, employee_id, year or date.today().year)


# --- teacher self-service ---


@router.get("/my/leaves", response_model=list[LeaveRequestOut])
def my_leaves(current_user: User = Depends(teacher_only), db: Session = Depends(get_db)):
    service = PayrollLeaveService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    return service.list_out(current_user.tenant_id, service.requests.list_filtered(current_user.tenant_id, teacher.id, None))


@router.post("/my/leaves", response_model=LeaveRequestOut, status_code=status.HTTP_201_CREATED)
def apply_my_leave(payload: MyLeaveRequestCreate, current_user: User = Depends(teacher_only), db: Session = Depends(get_db)):
    service = PayrollLeaveService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    employee = service.hr.resolve(current_user.tenant_id, "teacher", teacher.id)
    return service.to_out(service.apply(current_user.tenant_id, employee, payload, current_user.id))


@router.post("/my/leaves/{request_id}/cancel", response_model=LeaveRequestOut)
def cancel_my_leave(request_id: uuid.UUID, current_user: User = Depends(teacher_only), db: Session = Depends(get_db)):
    service = PayrollLeaveService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    return service.to_out(service.cancel_own(current_user.tenant_id, request_id, teacher.id))


@router.get("/my/leave-balances", response_model=list[LeaveBalanceOut])
def my_leave_balances(
    year: int | None = Query(default=None, ge=2000, le=2100),
    current_user: User = Depends(teacher_only),
    db: Session = Depends(get_db),
):
    service = PayrollLeaveService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    return service.balances_for(current_user.tenant_id, "teacher", teacher.id, year or date.today().year)


@router.get("/my/payslips", response_model=list[PayslipOut])
def my_payslips(current_user: User = Depends(teacher_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    return service.my_payslips(current_user.tenant_id, teacher.id)


@router.get("/my/payslips/{payslip_id}/pdf")
def my_payslip_pdf(payslip_id: uuid.UUID, current_user: User = Depends(teacher_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    teacher = service.hr.teacher_for_user(current_user.tenant_id, current_user.id)
    payslip, run = service.get_own_payslip(current_user.tenant_id, payslip_id, teacher.id)
    return _payslip_pdf(db, current_user.tenant_id, payslip, run)


# --- payroll runs ---


@router.post("/runs", response_model=PayrollRunOut, status_code=status.HTTP_201_CREATED)
def generate_payroll_run(payload: PayrollRunCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    run, skipped = service.generate(current_user.tenant_id, payload.period_month, payload.period_year,
                                    current_user.id, payload.notes)
    return service.run_out(run, skipped)


@router.get("/runs", response_model=list[PayrollRunOut])
def list_payroll_runs(current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollRunService(db).list_runs(current_user.tenant_id)


@router.get("/runs/{run_id}", response_model=PayrollRunDetail)
def get_payroll_run(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    return service.run_out(service.get_run(current_user.tenant_id, run_id), detail=True)


@router.delete("/runs/{run_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_payroll_run(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    PayrollRunService(db).delete_run(current_user.tenant_id, run_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/runs/{run_id}/approve", response_model=PayrollRunOut)
def approve_payroll_run(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    return service.run_out(service.approve(current_user.tenant_id, run_id, current_user.id))


@router.post("/runs/{run_id}/mark-paid", response_model=PayrollRunOut)
def mark_payroll_run_paid(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    return service.run_out(service.mark_paid(current_user.tenant_id, run_id))


@router.get("/runs/{run_id}/sheet.pdf")
def payroll_sheet_pdf(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    run = service.get_run(current_user.tenant_id, run_id)
    detail = service.run_out(run, detail=True).model_dump(mode="json")
    pdf = render_payroll_sheet(tenant_name=_tenant_name(db, current_user.tenant_id), month=run.period_month,
                               year=run.period_year, status=run.status.value, payslips=detail["payslips"])
    return _pdf(pdf, f"payroll-{run.period_year}-{run.period_month:02d}.pdf")


@router.get("/runs/{run_id}/sheet.xlsx")
def payroll_sheet_xlsx(run_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    run = service.get_run(current_user.tenant_id, run_id)
    payslips = service.run_out(run, detail=True).model_dump(mode="json")["payslips"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append([_tenant_name(db, current_user.tenant_id)])
    ws.append([f"Payroll Sheet - {period_label(run.period_month, run.period_year)}", run.status.value.title()])
    ws.append([])
    ws.append(SHEET_COLUMNS)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    rows = sheet_rows(payslips)
    for row in rows:
        ws.append(row)
    total = ["", "", "TOTAL", ""] + [round(sum(float(r[i]) for r in rows), 2) for i in range(4, len(SHEET_COLUMNS))]
    total[7] = ""
    ws.append(total)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=payroll-{run.period_year}-{run.period_month:02d}.xlsx"},
    )


# --- payslips ---


@router.get("/payslips/{payslip_id}", response_model=PayslipOut)
def get_payslip(payslip_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    service = PayrollRunService(db)
    payslip, run = service.get_payslip(current_user.tenant_id, payslip_id)
    return service.payslip_out(payslip, run)


@router.get("/payslips/{payslip_id}/pdf")
def payslip_pdf(payslip_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    payslip, run = PayrollRunService(db).get_payslip(current_user.tenant_id, payslip_id)
    return _payslip_pdf(db, current_user.tenant_id, payslip, run)


@router.post("/payslips/{payslip_id}/adjustments", response_model=PayslipOut, status_code=status.HTTP_201_CREATED)
def add_payslip_adjustment(
    payslip_id: uuid.UUID, payload: PayslipAdjustmentCreate,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    return PayrollRunService(db).add_adjustment(current_user.tenant_id, payslip_id, payload)


@router.delete("/payslips/{payslip_id}/adjustments/{adjustment_id}", response_model=PayslipOut)
def remove_payslip_adjustment(
    payslip_id: uuid.UUID, adjustment_id: uuid.UUID,
    current_user: User = Depends(admin_only), db: Session = Depends(get_db),
):
    return PayrollRunService(db).remove_adjustment(current_user.tenant_id, payslip_id, adjustment_id)


# --- advances ---


@router.get("/advances", response_model=list[SalaryAdvanceOut])
def list_advances(
    employee_type: EmployeeType | None = Query(default=None),
    employee_id: uuid.UUID | None = Query(default=None),
    status_filter: SalaryAdvanceStatus | None = Query(default=None, alias="status"),
    current_user: User = Depends(admin_only),
    db: Session = Depends(get_db),
):
    return PayrollRunService(db).list_advances(current_user.tenant_id, employee_type, employee_id, status_filter)


@router.post("/advances", response_model=SalaryAdvanceOut, status_code=status.HTTP_201_CREATED)
def create_advance(payload: SalaryAdvanceCreate, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    return PayrollRunService(db).create_advance(current_user.tenant_id, payload, current_user.id)


@router.delete("/advances/{advance_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_advance(advance_id: uuid.UUID, current_user: User = Depends(admin_only), db: Session = Depends(get_db)):
    PayrollRunService(db).delete_advance(current_user.tenant_id, advance_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)
