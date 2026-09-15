import io
import uuid

import openpyxl
from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.core.exceptions import DomainError
from app.schemas.student import StudentCreate
from app.schemas.student_admission_detail import StudentAdmissionDetailIn
from app.schemas.student_import import StudentImportResult, StudentImportRowError
from app.services.student_service import StudentService

SAMPLE_HEADERS = ["Full Name", "Email", "Password", "Roll Number", "Father Name", "Father Mobile"]


def build_sample_template() -> io.BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(SAMPLE_HEADERS)
    sheet.append(["Sam Student", "sam@example.com", "Password123!", "5A-01", "Ahmad Ali", "03001234567"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class StudentImportService:
    def __init__(self, db: Session):
        self.db = db
        self.student_service = StudentService(db)

    def import_from_excel(
        self,
        tenant_id: uuid.UUID,
        class_grade_id: uuid.UUID,
        section_id: uuid.UUID | None,
        file: UploadFile,
    ) -> StudentImportResult:
        self.student_service.validate_class_and_section(tenant_id, class_grade_id, section_id)

        contents = file.file.read()
        workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return StudentImportResult(created=0, failed=[])

        created = 0
        failed: list[StudentImportRowError] = []

        for index, row in enumerate(rows[1:], start=2):  # skip header row
            if row is None or all(cell is None for cell in row):
                continue
            full_name, email, password, roll_number, father_name, father_mobile = (list(row) + [None] * 6)[:6]
            try:
                if not full_name or not email or not password:
                    raise DomainError("Full Name, Email, and Password are required")
                payload = StudentCreate(
                    full_name=str(full_name),
                    email=str(email),
                    password=str(password),
                    class_grade_id=class_grade_id,
                    section_id=section_id,
                    roll_number=str(roll_number) if roll_number else None,
                    admission_detail=(
                        StudentAdmissionDetailIn(
                            father_name=str(father_name) if father_name else None,
                            father_mobile=str(father_mobile) if father_mobile else None,
                        )
                        if father_name or father_mobile
                        else None
                    ),
                )
                self.student_service.create_student(tenant_id, payload)
                created += 1
            except Exception as exc:  # noqa: BLE001 - per-row isolation, continue importing remaining rows
                self.db.rollback()
                message = exc.message if isinstance(exc, DomainError) else str(exc)
                failed.append(StudentImportRowError(row=index, error=message))

        return StudentImportResult(created=created, failed=failed)
